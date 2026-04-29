#!/usr/bin/env python3
"""生成 KiCad Symbol Web Renderer 实现原理 PPT。

该脚本只消费项目已有文档、研究结论和当前源码事实，不改变业务代码。
输出的 PPTX 使用 python-pptx 原生对象，便于后续继续编辑。
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from textwrap import dedent
from typing import Literal
import re
import zipfile

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DOC = ROOT / "docs" / "kicad-symbol-web-rendering.md"
OUTPUT_PPTX = ROOT / "docs" / "kicad-symbol-web-rendering-implementation.pptx"
PREVIEW_DIR = ROOT / "docs" / "ppt-previews"
OUTLINE_FILE = PREVIEW_DIR / "kicad-symbol-web-rendering-outline.md"
QA_FILE = PREVIEW_DIR / "kicad-symbol-web-rendering-qa.md"
MONTAGE_FILE = PREVIEW_DIR / "kicad-symbol-web-rendering-montage.png"

SLIDE_W_IN = 13.333333
SLIDE_H_IN = 7.5
PREVIEW_SCALE = 144
PREVIEW_W = int(SLIDE_W_IN * PREVIEW_SCALE)
PREVIEW_H = int(SLIDE_H_IN * PREVIEW_SCALE)

FONT_FACE = "PingFang SC"
BODY_FONT_FACE = "Microsoft YaHei"
MONO_FONT_FACE = "Menlo"
FONT_PATH = Path("/System/Library/Fonts/STHeiti Medium.ttc")
MONO_FONT_PATH = Path("/System/Library/Fonts/Menlo.ttc")

Color = str
Align = Literal["left", "center", "right"]


PALETTE = {
    "ink": "0F171A",
    "ink2": "182327",
    "paper": "F7F3EA",
    "paper2": "ECE4D5",
    "muted": "6C746F",
    "grid": "C8C3B6",
    "copper": "D67743",
    "gold": "F0B84A",
    "pcb": "0B6B5B",
    "silicon": "315C9F",
    "kicad": "840000",
    "green": "006464",
    "white": "FFFFFF",
    "soft_red": "F4D8CA",
    "soft_green": "D8ECE4",
    "soft_blue": "D8E4F4",
}


@dataclass
class PreviewElement:
    kind: str
    args: tuple
    kwargs: dict = field(default_factory=dict)


@dataclass
class SlideMeta:
    title: str
    subtitle: str = ""
    notes: str = ""


def rgb(color: Color) -> RGBColor:
    value = color.replace("#", "")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def pil_rgb(color: Color) -> tuple[int, int, int]:
    value = color.replace("#", "")
    return int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16)


def inch(value: float):
    return Inches(value)


def px(value: float) -> int:
    return int(round(value * PREVIEW_SCALE))


def load_font(size_pt: float, mono: bool = False) -> ImageFont.FreeTypeFont:
    path = MONO_FONT_PATH if mono and MONO_FONT_PATH.exists() else FONT_PATH
    if path.exists():
        return ImageFont.truetype(str(path), max(8, int(size_pt * 2)))
    return ImageFont.load_default()


def wrap_lines(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont, max_width: int) -> list[str]:
    lines: list[str] = []
    for paragraph in text.split("\n"):
        if not paragraph:
            lines.append("")
            continue
        current = ""
        for char in paragraph:
            candidate = current + char
            if current and draw.textlength(candidate, font=font) > max_width:
                lines.append(current)
                current = char
            else:
                current = candidate
        if current:
            lines.append(current)
    return lines


def draw_text_preview(
    draw: ImageDraw.ImageDraw,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    size: float,
    color: Color,
    align: Align = "left",
    mono: bool = False,
    line_gap: float = 1.16,
):
    font = load_font(size, mono=mono)
    max_width = px(w)
    lines = wrap_lines(draw, text, font, max_width)
    line_h = int(size * 2 * line_gap)
    y0 = px(y)
    max_y = px(y + h)
    for line in lines:
        if y0 + line_h > max_y:
            break
        text_w = draw.textlength(line, font=font)
        if align == "center":
            x0 = px(x) + (max_width - int(text_w)) // 2
        elif align == "right":
            x0 = px(x) + max_width - int(text_w)
        else:
            x0 = px(x)
        draw.text((x0, y0), line, font=font, fill=pil_rgb(color))
        y0 += line_h


class SlideBuilder:
    """把同一组几何描述同时写入 PPTX 和 PNG 预览。"""

    def __init__(self, prs: Presentation, meta: SlideMeta, bg: Color = PALETTE["paper"]):
        self.meta = meta
        self.slide = prs.slides.add_slide(prs.slide_layouts[6])
        self.slide.background.fill.solid()
        self.slide.background.fill.fore_color.rgb = rgb(bg)
        self.bg = bg
        self.elements: list[PreviewElement] = []
        self.notes = meta.notes
        if meta.notes:
            self.slide.notes_slide.notes_text_frame.text = meta.notes

    def rect(
        self,
        x: float,
        y: float,
        w: float,
        h: float,
        fill: Color,
        line: Color | None = None,
        line_width: float = 0.8,
        rounded: bool = False,
    ):
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
        shape = self.slide.shapes.add_shape(shape_type, inch(x), inch(y), inch(w), inch(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        if line:
            shape.line.color.rgb = rgb(line)
            shape.line.width = Pt(line_width)
        else:
            shape.line.fill.background()
        self.elements.append(PreviewElement("rect", (x, y, w, h, fill, line, line_width, rounded)))
        return shape

    def oval(self, x: float, y: float, w: float, h: float, fill: Color, line: Color | None = None, line_width: float = 1.0):
        shape = self.slide.shapes.add_shape(MSO_SHAPE.OVAL, inch(x), inch(y), inch(w), inch(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        if line:
            shape.line.color.rgb = rgb(line)
            shape.line.width = Pt(line_width)
        else:
            shape.line.fill.background()
        self.elements.append(PreviewElement("oval", (x, y, w, h, fill, line, line_width)))
        return shape

    def line(self, x1: float, y1: float, x2: float, y2: float, color: Color, width: float = 1.5):
        shape = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, inch(x1), inch(y1), inch(x2), inch(y2))
        shape.line.color.rgb = rgb(color)
        shape.line.width = Pt(width)
        self.elements.append(PreviewElement("line", (x1, y1, x2, y2, color, width)))
        return shape

    def text(
        self,
        x: float,
        y: float,
        w: float,
        h: float,
        text: str,
        size: float = 18,
        color: Color = PALETTE["ink"],
        bold: bool = False,
        align: Align = "left",
        valign: Literal["top", "mid", "bottom"] = "top",
        mono: bool = False,
        line_spacing: float = 1.05,
    ):
        # PowerPoint 对文本框高度较敏感。这里按字号给出最小高度，
        # 避免单行标题或标签在不同 Office 渲染器中被裁切。
        effective_h = max(h, (size / 72) * 1.45 * line_spacing)
        box = self.slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(effective_h))
        frame = box.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = 0
        frame.margin_right = 0
        frame.margin_top = 0
        frame.margin_bottom = 0
        frame.vertical_anchor = {
            "top": MSO_ANCHOR.TOP,
            "mid": MSO_ANCHOR.MIDDLE,
            "bottom": MSO_ANCHOR.BOTTOM,
        }[valign]
        align_value = {
            "left": PP_ALIGN.LEFT,
            "center": PP_ALIGN.CENTER,
            "right": PP_ALIGN.RIGHT,
        }[align]
        for idx, paragraph_text in enumerate(text.split("\n")):
            paragraph = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
            paragraph.alignment = align_value
            paragraph.line_spacing = line_spacing
            run = paragraph.add_run()
            run.text = paragraph_text
            run.font.name = MONO_FONT_FACE if mono else BODY_FONT_FACE
            run.font.size = Pt(size)
            run.font.color.rgb = rgb(color)
            run.font.bold = bold
        self.elements.append(PreviewElement("text", (x, y, w, effective_h, text, size, color, align, mono)))
        return box

    def title(self, title: str, subtitle: str = "", dark: bool = False):
        color = PALETTE["white"] if dark else PALETTE["ink"]
        sub_color = "C9D1CD" if dark else PALETTE["muted"]
        self.text(0.7, 0.42, 9.5, 0.52, title, 25, color, True)
        if subtitle:
            self.text(0.72, 0.96, 9.2, 0.32, subtitle, 11.5, sub_color)

    def footer(self, index: int, dark: bool = False):
        color = "B8C3BF" if dark else PALETTE["muted"]
        self.text(0.72, 7.03, 7.0, 0.2, "KiCad Symbol Web Renderer | 实现原理", 8.5, color)
        self.text(12.15, 7.03, 0.55, 0.2, f"{index:02d}", 8.5, color, align="right")

    def render_preview(self, path: Path):
        image = Image.new("RGB", (PREVIEW_W, PREVIEW_H), pil_rgb(self.bg))
        draw = ImageDraw.Draw(image)
        for element in self.elements:
            if element.kind == "rect":
                x, y, w, h, fill, line, line_width, rounded = element.args
                box = [px(x), px(y), px(x + w), px(y + h)]
                if rounded:
                    draw.rounded_rectangle(box, radius=px(0.12), fill=pil_rgb(fill), outline=pil_rgb(line) if line else None, width=max(1, int(line_width * 2)))
                else:
                    draw.rectangle(box, fill=pil_rgb(fill), outline=pil_rgb(line) if line else None, width=max(1, int(line_width * 2)))
            elif element.kind == "oval":
                x, y, w, h, fill, line, line_width = element.args
                draw.ellipse([px(x), px(y), px(x + w), px(y + h)], fill=pil_rgb(fill), outline=pil_rgb(line) if line else None, width=max(1, int(line_width * 2)))
            elif element.kind == "line":
                x1, y1, x2, y2, color, width = element.args
                draw.line([px(x1), px(y1), px(x2), px(y2)], fill=pil_rgb(color), width=max(1, int(width * 2)))
            elif element.kind == "text":
                x, y, w, h, text, size, color, align, mono = element.args
                draw_text_preview(draw, x, y, w, h, text, size, color, align=align, mono=mono)
        image.save(path)


def add_chip(slide: SlideBuilder, x: float, y: float, text: str, color: Color, w: float = 1.5):
    slide.rect(x, y, w, 0.34, color, rounded=True)
    slide.text(x + 0.09, y + 0.08, w - 0.18, 0.14, text, 8.5, PALETTE["white"], True, align="center")


def add_pipeline_node(slide: SlideBuilder, x: float, y: float, w: float, label: str, detail: str, fill: Color, stroke: Color):
    slide.rect(x, y, w, 0.72, fill, stroke, 1.2, rounded=True)
    slide.text(x + 0.12, y + 0.12, w - 0.24, 0.18, label, 10.8, PALETTE["ink"], True, align="center")
    slide.text(x + 0.12, y + 0.4, w - 0.24, 0.18, detail, 7.9, PALETTE["muted"], align="center")


def add_section_marker(slide: SlideBuilder, x: float, y: float, number: str, label: str, color: Color):
    slide.oval(x, y, 0.46, 0.46, color)
    slide.text(x, y + 0.12, 0.46, 0.12, number, 8.5, PALETTE["white"], True, align="center")
    slide.text(x + 0.58, y + 0.03, 2.2, 0.23, label, 11.8, PALETTE["ink"], True)


def add_manual_table(slide: SlideBuilder, x: float, y: float, col_w: list[float], row_h: list[float], data: list[list[str]], header_fill: Color):
    yy = y
    for r, height in enumerate(row_h):
        xx = x
        for c, width in enumerate(col_w):
            fill = header_fill if r == 0 else (PALETTE["paper"] if r % 2 else "F2EBDD")
            slide.rect(xx, yy, width, height, fill, PALETTE["grid"], 0.55)
            text_color = PALETTE["white"] if r == 0 else PALETTE["ink"]
            size = 8.2 if r else 8.8
            slide.text(xx + 0.08, yy + 0.08, width - 0.16, height - 0.12, data[r][c], size, text_color, r == 0)
            xx += width
        yy += height


def create_deck() -> tuple[Presentation, list[SlideBuilder]]:
    if not SOURCE_DOC.exists():
        raise FileNotFoundError(f"缺少源文档: {SOURCE_DOC}")

    prs = Presentation()
    prs.slide_width = inch(SLIDE_W_IN)
    prs.slide_height = inch(SLIDE_H_IN)
    props = prs.core_properties
    props.title = "KiCad 符号 Web 渲染实现原理"
    props.subject = "从 KiCad S-expression 到 Symbol IR、SVG/PixiJS 渲染和可回写导出"
    props.author = "Codex"
    props.keywords = "KiCad, EDA, Symbol IR, S-expression, PixiJS, SVG"

    slides: list[SlideBuilder] = []

    # 01 封面：用 S-expression 代码片段和坐标网格作为主题视觉。
    s = SlideBuilder(prs, SlideMeta("封面"), bg=PALETTE["ink"])
    slides.append(s)
    for i in range(0, 18):
        x = 0.5 + i * 0.72
        s.line(x, 0.0, x, 7.5, "253135", 0.5)
    for j in range(0, 11):
        y = 0.35 + j * 0.62
        s.line(0.0, y, 13.33, y, "253135", 0.5)
    s.rect(7.55, 0.65, 4.85, 5.85, "111F23", "385048", 1.0, rounded=True)
    code = dedent(
        """
        (kicad_symbol_lib
          (version 20251024)
          (symbol "TLP250"
            (property "Reference" "U")
            (symbol "TLP250_1_1"
              (pin power_in line
                (at 10.16 7.62 180)
                (length 2.54)))))
        """
    ).strip()
    s.text(7.9, 1.05, 4.15, 3.6, code, 10.2, "D6E8DE", mono=True, line_spacing=1.05)
    s.rect(7.9, 4.9, 3.45, 0.18, PALETTE["kicad"])
    s.text(0.75, 0.72, 5.9, 0.35, "实现原理详解", 16, PALETTE["gold"], True)
    s.text(0.72, 1.28, 6.35, 1.15, "KiCad 符号\nWeb 端渲染", 43, PALETTE["white"], True, line_spacing=0.92)
    s.text(0.78, 3.0, 6.35, 0.72, "从 .kicad_sym S-expression 到统一 IR、SVG 基准、\nPixiJS 交互投影和可回写导出", 15.8, "D9E0DC")
    add_chip(s, 0.8, 4.0, "TypeScript IR", PALETTE["pcb"], 1.75)
    add_chip(s, 2.75, 4.0, "SVG 基准", PALETTE["kicad"], 1.35)
    add_chip(s, 4.3, 4.0, "PixiJS 投影", PALETTE["silicon"], 1.62)
    s.text(0.78, 6.72, 5.2, 0.22, "基于 docs/kicad-symbol-web-rendering.md 细化生成", 9.5, "AAB8B3")

    # 02 目标与边界。
    s = SlideBuilder(prs, SlideMeta("目标与边界"))
    slides.append(s)
    s.title("核心目标不是“画出来”，而是 KiCad 符号互操作")
    s.text(0.85, 1.55, 5.2, 0.85, "任意受支持的 `.kicad_sym` 都能在 Web 端高保真渲染，并且 Web 创建的符号能无语义丢失地回写为 KiCad 可识别文件。", 19.5, PALETTE["ink"], True)
    s.rect(7.0, 1.15, 4.6, 1.28, PALETTE["ink"], rounded=True)
    s.text(7.3, 1.42, 3.9, 0.28, "KiCad Symbol IR", 21, PALETTE["white"], True, align="center")
    s.text(7.28, 1.92, 3.9, 0.22, "唯一业务真相", 12.8, PALETTE["gold"], True, align="center")
    pairs = [
        ("格式兼容", ".kicad_sym 按 KiCad S-expression 语义解析和生成"),
        ("数值精度", "坐标统一按 mm 处理，导出最多四位小数"),
        ("渲染一致性", "SVG 对齐 KiCad CLI，PixiJS 追随同一 IR"),
        ("安全防线", "外部符号库、文本、SVG 输出和大文件都要受限"),
    ]
    for idx, (label, detail) in enumerate(pairs):
        y = 3.05 + idx * 0.68
        s.oval(0.95, y, 0.25, 0.25, [PALETTE["kicad"], PALETTE["pcb"], PALETTE["silicon"], PALETTE["copper"]][idx])
        s.text(1.35, y - 0.02, 1.3, 0.23, label, 12.8, PALETTE["ink"], True)
        s.text(2.65, y - 0.02, 7.5, 0.25, detail, 12.2, PALETTE["muted"])
    s.rect(8.0, 4.35, 2.55, 0.65, PALETTE["soft_green"], PALETTE["pcb"], 1.0)
    s.text(8.15, 4.56, 2.25, 0.15, "第一阶段：符号库本体", 9.4, PALETTE["pcb"], True, align="center")
    s.text(7.38, 5.42, 3.95, 0.45, "暂不扩展到完整原理图编辑、PCB 编辑或仿真。", 13.5, PALETTE["ink"], True, align="center")
    s.footer(2)

    # 03 关键问题分解。
    s = SlideBuilder(prs, SlideMeta("关键问题"))
    slides.append(s)
    s.title("Web 端渲染 KiCad 符号需要同时解决四类正确性")
    problems = [
        ("文件语法", "S-expression 嵌套、字符串转义、未知 token、source span"),
        ("领域语义", "symbol、unit、property、graphic、pin、effects 的强类型表达"),
        ("视觉投影", "mm 坐标、Y 轴翻转、stroke font、pin 样式、bounds"),
        ("导出闭环", "Web 编辑结果必须能回写 KiCad，不能从画布反推语义"),
    ]
    for idx, (label, detail) in enumerate(problems):
        x = 0.85 + (idx % 2) * 5.85
        y = 1.55 + (idx // 2) * 2.25
        color = [PALETTE["kicad"], PALETTE["pcb"], PALETTE["silicon"], PALETTE["copper"]][idx]
        s.line(x, y + 0.12, x, y + 1.55, color, 4)
        s.text(x + 0.25, y, 3.8, 0.3, label, 18.5, PALETTE["ink"], True)
        s.text(x + 0.25, y + 0.52, 4.55, 0.8, detail, 13.5, PALETTE["muted"])
        s.text(x + 4.75, y + 0.02, 0.4, 0.3, f"{idx + 1}", 18, color, True, align="right")
    s.text(1.05, 6.35, 10.6, 0.32, "结论：Renderer 只能是投影层，不能成为文件格式、业务语义或导出逻辑的源头。", 16, PALETTE["ink"], True, align="center")
    s.footer(3)

    # 04 总体链路图。
    s = SlideBuilder(prs, SlideMeta("总体流程"))
    slides.append(s)
    s.title(".kicad_sym 到浏览器画布的主链路")
    nodes = [
        (0.75, 1.55, ".kicad_sym", "KiCad 文件"),
        (2.35, 1.55, "Tokenizer", "token + raw"),
        (3.95, 1.55, "CST", "顺序 + span"),
        (5.55, 1.55, "AST", "KiCad 语义"),
        (7.15, 1.55, "IR", "唯一真相"),
        (8.75, 1.55, "Layout", "bounds + pin"),
        (10.35, 1.55, "PixiJS", "Canvas 预览"),
    ]
    for i, (x, y, label, detail) in enumerate(nodes):
        add_pipeline_node(s, x, y, 1.25, label, detail, ["F4E0D4", "E7E0D0", "DFEADC", "D8E4F4", "D8ECE4", "F3E8C4", "F1D8D2"][i], [PALETTE["copper"], PALETTE["muted"], PALETTE["pcb"], PALETTE["silicon"], PALETTE["pcb"], PALETTE["gold"], PALETTE["kicad"]][i])
        if i < len(nodes) - 1:
            s.line(x + 1.25, y + 0.36, nodes[i + 1][0], y + 0.36, PALETTE["grid"], 1.2)
    s.line(7.75, 2.27, 7.75, 4.9, PALETTE["pcb"], 1.5)
    add_pipeline_node(s, 5.9, 4.55, 1.4, "SVG", "黄金基准", "D8E4F4", PALETTE["silicon"])
    add_pipeline_node(s, 7.05, 5.45, 1.55, "Editor", "拖拽 / Inspector", "D8ECE4", PALETTE["pcb"])
    add_pipeline_node(s, 8.85, 4.55, 1.55, "Exporter", "物化 AST/CST", "F4E0D4", PALETTE["copper"])
    s.line(7.75, 2.27, 6.6, 4.55, PALETTE["grid"], 1.1)
    s.line(7.75, 2.27, 7.82, 5.45, PALETTE["grid"], 1.1)
    s.line(7.75, 2.27, 9.62, 4.55, PALETTE["grid"], 1.1)
    s.text(0.95, 3.55, 3.8, 1.1, "关键思想：只允许数据从文件语义流入 IR，再由 IR 派生渲染、编辑和导出。SVG、PixiJS、DOM 或 X6 都不能反向成为业务数据源。", 14.8, PALETTE["ink"], True)
    s.footer(4)

    # 05 分层职责。
    s = SlideBuilder(prs, SlideMeta("分层职责"))
    slides.append(s)
    s.title("分层职责：每一层只做自己的事")
    table = [
        ["层级", "当前代码", "职责", "禁止越界"],
        ["Tokenizer", "src/sexpr/tokenizer.ts", "识别括号、字符串、数字、符号，处理资源上限", "不理解 KiCad 语义"],
        ["Parser / Printer", "src/sexpr/parser.ts", "生成 CST，保留 raw atom 和 source span", "不做布局或单位换算"],
        ["KiCad AST", "src/kicad/parser.ts", "映射 symbol、property、graphic、pin、effects", "不依赖 DOM 或 PixiJS"],
        ["Symbol IR", "src/kicad/ir.ts", "稳定 ID、编辑命令、renderer 无关模型", "不存渲染器状态"],
        ["Layout", "src/web/render/symbol-layout.ts", "汇总图元、pin、标签、ref 和 bounds", "不修改 IR"],
        ["PixiJS", "SymbolCanvas.ts / App.vue", "场景投影、命中热区、Inspector、导出", "不成为持久化模型"],
    ]
    add_manual_table(s, 0.55, 1.35, [1.4, 2.55, 4.5, 3.4], [0.45] + [0.66] * 6, table, PALETTE["ink"])
    s.text(0.78, 6.45, 11.8, 0.28, "这种边界让 SVG renderer、PixiJS renderer、Web editor 和 exporter 可以共享 IR，又互不污染职责。", 13.8, PALETTE["ink"], True, align="center")
    s.footer(5)

    # 06 S-expression 解析。
    s = SlideBuilder(prs, SlideMeta("S-expression 解析"))
    slides.append(s)
    s.title("第一层：安全读取 S-expression，而不是用正则猜结构")
    s.rect(0.82, 1.35, 3.75, 4.6, PALETTE["ink"], rounded=True)
    code = dedent(
        """
        token kinds:
          open  "("
          close ")"
          atom  string | number | symbol

        string escapes:
          \\"  \\\\  \\n  \\r  \\t

        CST keeps:
          raw
          value
          source span
        """
    ).strip()
    s.text(1.12, 1.75, 3.15, 3.65, code, 11.5, "D8ECE4", mono=True)
    limits = [("5MB", "默认输入上限"), ("128", "最大嵌套层级"), ("200000", "最大 token 数")]
    for idx, (num, label) in enumerate(limits):
        x = 5.35 + idx * 2.25
        s.text(x, 1.65, 1.8, 0.42, num, 31, [PALETTE["kicad"], PALETTE["pcb"], PALETTE["silicon"]][idx], True, align="center")
        s.text(x, 2.18, 1.8, 0.24, label, 10.8, PALETTE["muted"], align="center")
    points = [
        ("资源上限前置", "避免超大文件、深层嵌套和海量 token 造成浏览器或 Node 资源耗尽。"),
        ("raw 与 span 保留", "为错误定位、可控打印和未知字段 round-trip 留出基础。"),
        ("语法层保持纯粹", "tokenizer 只处理通用 S-expression，不提前混入 KiCad 图形规则。"),
    ]
    for idx, (head, body) in enumerate(points):
        y = 3.25 + idx * 0.78
        s.oval(5.4, y, 0.23, 0.23, PALETTE["copper"])
        s.text(5.78, y - 0.05, 2.25, 0.22, head, 13.3, PALETTE["ink"], True)
        s.text(7.68, y - 0.05, 3.7, 0.28, body, 11.5, PALETTE["muted"])
    s.footer(6)

    # 07 KiCad AST。
    s = SlideBuilder(prs, SlideMeta("KiCad AST"))
    slides.append(s)
    s.title("第二层：从 CST 映射成 KiCad Symbol AST")
    tree_x = 0.95
    tree_y = 1.45
    tree = [
        ("kicad_symbol_lib", PALETTE["ink"], 0),
        ("header: version / generator", PALETTE["copper"], 1),
        ("symbol: TLP250", PALETTE["pcb"], 1),
        ("property: Reference / Value / ...", PALETTE["silicon"], 2),
        ("graphic: rectangle / circle / arc / bezier / text", PALETTE["kicad"], 2),
        ("pin: electrical type / style / at / length", PALETTE["gold"], 2),
        ("unknownEntries", PALETTE["muted"], 1),
    ]
    for idx, (label, color, depth) in enumerate(tree):
        x = tree_x + depth * 0.45
        y = tree_y + idx * 0.62
        s.line(tree_x + 0.08 + (depth - 1) * 0.45 if depth else x, y + 0.17, x, y + 0.17, PALETTE["grid"], 1)
        s.oval(x, y + 0.08, 0.16, 0.16, color)
        s.text(x + 0.28, y, 4.8, 0.22, label, 12.2, PALETTE["ink"], depth == 0)
    s.rect(7.05, 1.48, 4.5, 3.3, "F3E8C4", PALETTE["gold"], 1.0, rounded=True)
    s.text(7.32, 1.83, 3.95, 0.28, "AST 的价值", 18, PALETTE["ink"], True, align="center")
    s.text(7.35, 2.42, 3.85, 1.85, "把文件里的 KiCad 语义明确化：属性、图元、引脚、字体效果、线宽、填充、显示策略都变成强类型结构。未知字段挂在 unknownEntries，当前版本暂不理解也不静默丢弃。", 14, PALETTE["ink"])
    s.text(1.1, 6.1, 10.7, 0.36, "解析器的边界：只解释 KiCad 文件语义，不做像素坐标、画布对象或 Web 编辑状态。", 15, PALETTE["kicad"], True, align="center")
    s.footer(7)

    # 08 IR 设计。
    s = SlideBuilder(prs, SlideMeta("Symbol IR"))
    slides.append(s)
    s.title("第三层：Symbol IR 让所有投影引用同一份业务对象")
    s.rect(0.75, 1.52, 4.35, 1.22, PALETTE["soft_green"], PALETTE["pcb"], 1.1, rounded=True)
    s.text(1.05, 1.82, 3.78, 0.25, "symbol:TLP250", 18.5, PALETTE["pcb"], True, align="center", mono=True)
    examples = [
        "symbol:TLP250/property:Reference",
        "symbol:TLP250/unit:TLP250_1_1",
        "symbol:TLP250/unit:TLP250_1_1/pin:8:VCC",
        "symbol:TLP250/unit:TLP250_0_1/graphic:rectangle:1",
    ]
    for idx, text in enumerate(examples):
        y = 3.15 + idx * 0.55
        s.text(0.92, y, 5.5, 0.18, text, 10.4, PALETTE["ink"], mono=True)
    s.line(6.6, 1.35, 6.6, 5.75, PALETTE["grid"], 1.1)
    principles = [
        ("稳定 ID", "选中态、拖拽、测试断言、SVG 和 PixiJS 都能定位同一对象。"),
        ("unit 规范化", "顶层图元也提升为虚拟 unit，renderer 只遍历一套结构。"),
        ("renderer 无关", "IR 不保存 PixiJS DisplayObject、SVG 节点或 DOM 状态。"),
        ("语义可回写", "编辑命令改 IR，再物化回 AST/CST 并序列化为 KiCad 文件。"),
    ]
    for idx, (head, body) in enumerate(principles):
        y = 1.52 + idx * 0.95
        s.oval(7.02, y + 0.02, 0.32, 0.32, [PALETTE["pcb"], PALETTE["copper"], PALETTE["silicon"], PALETTE["kicad"]][idx])
        s.text(7.55, y, 1.55, 0.25, head, 13.5, PALETTE["ink"], True)
        s.text(9.05, y, 3.05, 0.35, body, 11.2, PALETTE["muted"])
    s.footer(8)

    # 09 布局与几何。
    s = SlideBuilder(prs, SlideMeta("布局与几何"))
    slides.append(s)
    s.title("Layout 层：把 IR 整理成可渲染输入，但不修改 IR")
    s.rect(0.82, 1.35, 5.1, 4.55, "F1EAD9", "D3C7B4", 0.8, rounded=True)
    # 简化符号图形示意。
    s.rect(2.15, 2.15, 1.75, 1.9, "FFF6C7", PALETTE["kicad"], 2.0)
    for idx, (x1, y1, x2, y2) in enumerate([(1.25, 2.55, 2.15, 2.55), (1.25, 3.15, 2.15, 3.15), (3.9, 2.55, 4.85, 2.55), (3.9, 3.15, 4.85, 3.15)]):
        s.line(x1, y1, x2, y2, PALETTE["kicad"], 2)
        s.text((x1 + x2) / 2 - 0.12, y1 - 0.35, 0.25, 0.16, str(idx + 1), 8, PALETTE["kicad"], True, align="center")
    s.line(1.05, 4.55, 5.45, 4.55, PALETTE["pcb"], 1.2)
    s.line(1.05, 4.55, 1.05, 1.75, PALETTE["silicon"], 1.2)
    s.text(1.08, 4.7, 1.2, 0.18, "bounds", 9.5, PALETTE["pcb"], True)
    details = [
        ("图元收集", "units.flatMap(graphics)"),
        ("pin 收集", "units.flatMap(pins)"),
        ("标签收集", "properties"),
        ("边界计算", "rectangle / pts / circle / arc / pinEndPoint / text"),
    ]
    for idx, (head, body) in enumerate(details):
        y = 1.55 + idx * 0.85
        s.text(6.75, y, 1.5, 0.22, head, 13, PALETTE["ink"], True)
        s.text(8.2, y, 3.8, 0.22, body, 11, PALETTE["muted"], mono=idx > 0)
    s.text(6.75, 5.3, 5.1, 0.32, "后续应抽象独立 geometry resolver，让 SVG 与 PixiJS 共用 bounds、pin 端点和圆弧求解。", 14, PALETTE["kicad"], True)
    s.footer(9)

    # 10 坐标变换。
    s = SlideBuilder(prs, SlideMeta("坐标变换"))
    slides.append(s)
    s.title("坐标变换：KiCad 世界坐标到 Canvas 屏幕坐标")
    s.rect(0.95, 1.35, 5.1, 4.35, "F6F0E2", "D3C7B4", 0.8)
    s.line(1.45, 4.7, 5.4, 4.7, PALETTE["pcb"], 2)
    s.line(1.45, 4.7, 1.45, 1.8, PALETTE["silicon"], 2)
    s.text(5.18, 4.86, 0.35, 0.18, "X", 10, PALETTE["pcb"], True)
    s.text(1.12, 1.72, 0.35, 0.18, "Y", 10, PALETTE["silicon"], True)
    s.oval(3.35, 3.1, 0.18, 0.18, PALETTE["kicad"])
    s.text(3.58, 3.04, 1.65, 0.16, "KiCad mm 点", 9, PALETTE["ink"])
    s.rect(7.1, 1.35, 4.9, 4.35, "EDF2F7", "C7D1DE", 0.8)
    s.line(7.6, 2.1, 11.4, 2.1, PALETTE["pcb"], 2)
    s.line(7.6, 2.1, 7.6, 5.0, PALETTE["silicon"], 2)
    s.text(11.22, 2.27, 0.35, 0.18, "X", 10, PALETTE["pcb"], True)
    s.text(7.27, 5.05, 0.35, 0.18, "Y", 10, PALETTE["silicon"], True)
    s.oval(9.55, 3.7, 0.18, 0.18, PALETTE["kicad"])
    s.text(9.78, 3.65, 1.75, 0.16, "Canvas px 点", 9, PALETTE["ink"])
    s.line(6.28, 3.45, 6.85, 3.45, PALETTE["grid"], 2)
    formulas = [
        "scale = min((width - padding*2)/boundsW, (height - padding*2)/boundsH, 46px/mm)",
        "x = width/2 + (point.x - centerX) * scale",
        "y = height/2 - (point.y - centerY) * scale",
    ]
    for idx, formula in enumerate(formulas):
        s.text(1.05, 6.05 + idx * 0.32, 11.4, 0.18, formula, 9.7, PALETTE["ink"], mono=True)
    s.footer(10)

    # 11 PixiJS 投影。
    s = SlideBuilder(prs, SlideMeta("PixiJS 投影"))
    slides.append(s)
    s.title("PixiJS 负责交互投影，业务状态仍然只在 Symbol IR")
    stages = [
        ("Vue 组件挂载", "创建 Pixi Application"),
        ("IR 变化或尺寸变化", "重新生成 layout"),
        ("清空 stage", "从 IR 再投影一次"),
        ("绘制 + 命中", "graphic / pin / label\n+ overlay"),
    ]
    for idx, (head, body) in enumerate(stages):
        x = 0.9 + idx * 2.95
        s.oval(x, 1.55, 0.55, 0.55, [PALETTE["pcb"], PALETTE["silicon"], PALETTE["copper"], PALETTE["kicad"]][idx])
        s.text(x + 0.12, 1.72, 0.3, 0.12, str(idx + 1), 9, PALETTE["white"], True, align="center")
        s.text(x, 2.28, 2.25, 0.24, head, 13.2, PALETTE["ink"], True, align="center")
        s.text(x - 0.18, 2.72, 2.55, 0.36, body, 10.8, PALETTE["muted"], align="center")
        if idx < 3:
            s.line(x + 0.58, 1.82, x + 2.68, 1.82, PALETTE["grid"], 1.5)
    s.rect(1.2, 4.15, 4.9, 1.05, PALETTE["soft_blue"], PALETTE["silicon"], 1.0, rounded=True)
    s.text(1.45, 4.45, 4.4, 0.22, "kicad-svg 模式：更接近 KiCad 导出预览", 13.2, PALETTE["silicon"], True, align="center")
    s.rect(7.15, 4.15, 4.9, 1.05, PALETTE["soft_green"], PALETTE["pcb"], 1.0, rounded=True)
    s.text(7.4, 4.45, 4.4, 0.22, "editor 模式：网格、隐藏项、电气类型、拖拽热区", 13.2, PALETTE["pcb"], True, align="center")
    s.text(1.35, 5.85, 10.7, 0.28, "模式差异和命中热区都属于显示策略；拖拽只上报 ref + deltaMm，由 Vue 调用 IR 命令更新业务状态。", 15, PALETTE["kicad"], True, align="center")
    s.footer(11)

    # 12 SVG 与 PixiJS 双渲染。
    s = SlideBuilder(prs, SlideMeta("双渲染策略"))
    slides.append(s)
    s.title("为什么 SVG 和 PixiJS 都必须从 IR 出发")
    s.rect(5.35, 1.42, 2.65, 0.9, PALETTE["ink"], rounded=True)
    s.text(5.72, 1.75, 1.9, 0.18, "Symbol IR", 18, PALETTE["white"], True, align="center")
    branches = [
        (1.05, 3.28, "SVG 精确渲染器", "确定性输出、结构快照、KiCad CLI 黄金基准", PALETTE["silicon"]),
        (5.15, 3.28, "PixiJS 交互投影", "缩放、平移、选中、hover、缓存和批量实例性能", PALETTE["pcb"]),
        (9.25, 3.28, "Exporter / Editor", "命令修改 IR，物化 AST/CST 后导出 `.kicad_sym`", PALETTE["copper"]),
    ]
    for x, y, head, body, color in branches:
        s.line(6.68, 2.32, x + 1.45, y, PALETTE["grid"], 1.3)
        s.rect(x, y, 3.0, 1.42, "FFFFFF", color, 1.2, rounded=True)
        s.text(x + 0.22, y + 0.22, 2.55, 0.22, head, 13.3, color, True, align="center")
        s.text(x + 0.22, y + 0.62, 2.55, 0.4, body, 10.6, PALETTE["muted"], align="center")
    s.text(0.95, 6.0, 11.3, 0.48, "收益：解析只做一次，编辑只改 IR，导出只从 IR/AST/CST 派生，测试能精确定位是解析问题、编辑物化问题还是渲染投影问题。", 14.3, PALETTE["ink"], True, align="center")
    s.footer(12)

    # 13 导出与 round-trip。
    s = SlideBuilder(prs, SlideMeta("导出闭环"))
    slides.append(s)
    s.title("导出闭环：编辑后的 IR 已可物化回 .kicad_sym")
    chain = [
        ("IR", "编辑后的唯一真相", PALETTE["pcb"]),
        ("AST", "KiCad 语义树", PALETTE["silicon"]),
        ("CST", "顺序 / raw / span", PALETTE["copper"]),
        (".kicad_sym", "KiCad 可打开", PALETTE["kicad"]),
    ]
    for idx, (head, sub, color) in enumerate(chain):
        x = 1.1 + idx * 3.0
        s.rect(x, 1.7, 2.05, 0.85, "FFFFFF", color, 1.3, rounded=True)
        s.text(x + 0.2, 1.94, 1.65, 0.18, head, 17, color, True, align="center", mono=head.startswith("."))
        s.text(x + 0.2, 2.32, 1.65, 0.16, sub, 8.8, PALETTE["muted"], align="center")
        if idx < 3:
            s.line(x + 2.05, 2.12, x + 2.82, 2.12, PALETTE["grid"], 1.4)
    s.rect(0.95, 3.35, 5.35, 1.6, PALETTE["soft_green"], PALETTE["pcb"], 1.0, rounded=True)
    s.text(1.25, 3.78, 4.75, 0.22, "当前 materialize 会克隆原始 CST，再同步受支持字段", 13.4, PALETTE["pcb"], True, align="center")
    s.text(1.3, 4.25, 4.65, 0.26, "覆盖：坐标、文本、线宽、类型、长度、隐藏。", 11.5, PALETTE["ink"], align="center")
    s.rect(7.05, 3.35, 5.1, 1.6, PALETTE["soft_red"], PALETTE["kicad"], 1.0, rounded=True)
    s.text(7.35, 3.78, 4.55, 0.22, "导出 guardrail 仍需覆盖更复杂编辑", 13.4, PALETTE["kicad"], True, align="center")
    s.text(7.4, 4.25, 4.45, 0.26, "复杂编辑后续统一进入导出 guardrail。", 11.5, PALETTE["ink"], align="center")
    s.footer(13)

    # 14 安全与兼容边界。
    s = SlideBuilder(prs, SlideMeta("安全与兼容"))
    slides.append(s)
    s.title("安全与兼容：外部符号库默认不可信")
    rows = [
        ["风险", "防线", "落实位置"],
        ["资源耗尽", "输入字节数、token 数、嵌套深度限制", "tokenizeSExpr"],
        ["语义丢失", "未知 token 保留到 unknownEntries 和 CST source", "parser / serializer"],
        ["注入风险", "文本进入 SVG 或 DOM 前必须转义", "后续 SVG/DOM 投影"],
        ["不可导出图元", "编辑工具只开放 KiCad 可表达图元集", "后续 editor validator"],
        ["性能崩溃", "Web Worker 解析、PixiJS RenderTexture 缓存、视口裁剪", "后续大库阶段"],
    ]
    add_manual_table(s, 0.82, 1.32, [2.0, 6.0, 3.6], [0.5] + [0.78] * 5, rows, PALETTE["kicad"])
    s.text(1.0, 6.28, 11.1, 0.3, "原则：接受外部文件时先限制资源，再保留未知结构，最后才进入业务语义和渲染投影。", 14.5, PALETTE["ink"], True, align="center")
    s.footer(14)

    # 15 验证体系。
    s = SlideBuilder(prs, SlideMeta("验证体系"))
    slides.append(s)
    s.title("验证体系：从结构正确到视觉一致，再到 KiCad 自身校验")
    gates = [
        ("当前", "Vitest", "解析/打印、异常输入、TLP250 AST/IR、稳定 ID、round-trip、未知字段、拖拽与 Inspector 导出", PALETTE["pcb"]),
        ("Phase 2", "SVG 快照", "同一 IR 输出稳定 SVG，结构变化可审查", PALETTE["silicon"]),
        ("Phase 3", "KiCad CLI", "kicad-cli sym export svg 生成黄金基准", PALETTE["copper"]),
        ("Phase 4", "浏览器视觉回归", "Playwright 截图或 raster diff 验证 PixiJS 追随 SVG 基准", PALETTE["kicad"]),
    ]
    for idx, (phase, name, body, color) in enumerate(gates):
        y = 1.4 + idx * 1.13
        s.oval(0.95, y + 0.1, 0.38, 0.38, color)
        s.text(1.55, y + 0.03, 1.2, 0.22, phase, 11.5, color, True)
        s.text(2.85, y + 0.03, 2.0, 0.22, name, 14.2, PALETTE["ink"], True)
        s.text(4.85, y + 0.03, 6.85, 0.34, body, 11.5, PALETTE["muted"])
        if idx < 3:
            s.line(1.14, y + 0.48, 1.14, y + 1.13, PALETTE["grid"], 1.2)
    s.rect(8.8, 5.75, 2.5, 0.55, PALETTE["ink"], rounded=True)
    s.text(9.05, 5.95, 2.0, 0.13, "测试即互操作契约", 9.5, PALETTE["white"], True, align="center")
    s.footer(15)

    # 16 当前实现状态。
    s = SlideBuilder(prs, SlideMeta("当前实现"))
    slides.append(s)
    s.title("当前薄切片：TLP250 已跑通预览、编辑与导出回写")
    flow = ["tlp250.kicad_sym", "parseKicadSymbolLibrary", "toSymbolLibraryIR", "createSymbolLayout", "SymbolCanvas", "Canvas + 导出"]
    for idx, label in enumerate(flow):
        x = 0.65 + idx * 2.05
        s.rect(x, 1.55, 1.55, 0.72, "FFFFFF", [PALETTE["kicad"], PALETTE["copper"], PALETTE["pcb"], PALETTE["gold"], PALETTE["silicon"], PALETTE["ink"]][idx], 1.0, rounded=True)
        s.text(x + 0.1, 1.8, 1.35, 0.14, label, 7.8 if idx > 0 else 8.6, PALETTE["ink"], True, align="center")
        if idx < len(flow) - 1:
            s.line(x + 1.55, 1.91, x + 2.0, 1.91, PALETTE["grid"], 1.0)
    metrics = [("1", "symbol"), ("7", "properties"), ("8", "pins"), (">20", "graphics"), ("2", "render modes")]
    for idx, (num, label) in enumerate(metrics):
        x = 1.0 + idx * 2.35
        s.text(x, 3.1, 1.5, 0.45, num, 31, [PALETTE["kicad"], PALETTE["pcb"], PALETTE["silicon"], PALETTE["copper"], PALETTE["gold"]][idx], True, align="center")
        s.text(x, 3.7, 1.5, 0.18, label, 10.2, PALETTE["muted"], align="center")
    s.rect(1.25, 4.65, 10.4, 0.95, PALETTE["soft_blue"], PALETTE["silicon"], 1.0, rounded=True)
    s.text(1.55, 4.96, 9.8, 0.2, "Vue 持有 editableIR；PixiJS 只上报选择/拖拽；导出前物化回 KiCad AST/CST。", 12.8, PALETTE["ink"], True, align="center")
    s.footer(16)

    # 17 演进路线。
    s = SlideBuilder(prs, SlideMeta("演进路线"))
    slides.append(s)
    s.title("推荐演进路线：先精度闭环，再扩展编辑深度")
    roadmap = [
        ("1", "geometry resolver", "统一 bounds、pin 端点、圆弧和坐标变换"),
        ("2", "SVG renderer", "建立可与 KiCad CLI 对照的精确静态渲染"),
        ("3", "CLI validation harness", "固定 TLP250 与更多 fixture 的黄金输出"),
        ("4", "PixiJS caching", "静态图元 RenderTexture、交互层分离、culling"),
        ("5", "Editing expansion", "新增/删除、撤销重做、约束校验、导出 guardrail"),
    ]
    for idx, (num, head, body) in enumerate(roadmap):
        x = 0.95 + idx * 2.42
        color = [PALETTE["pcb"], PALETTE["silicon"], PALETTE["copper"], PALETTE["kicad"], PALETTE["gold"]][idx]
        s.oval(x, 1.65, 0.55, 0.55, color)
        s.text(x + 0.16, 1.82, 0.22, 0.12, num, 9, PALETTE["white"], True, align="center")
        s.line(x + 0.55, 1.92, x + 2.1, 1.92, PALETTE["grid"], 1.2)
        s.text(x - 0.1, 2.55, 1.95, 0.28, head, 11.8, PALETTE["ink"], True, align="center", mono=True)
        s.text(x - 0.08, 3.1, 1.9, 0.62, body, 10.2, PALETTE["muted"], align="center")
    s.text(1.1, 5.25, 10.9, 0.62, "取舍：不要过早进入完整原理图编辑或多人协同。当前核心风险仍是格式、语义、渲染精度和编辑导出闭环。", 16, PALETTE["kicad"], True, align="center")
    s.footer(17)

    # 18 结尾原则。
    s = SlideBuilder(prs, SlideMeta("设计原则"), bg=PALETTE["ink"])
    slides.append(s)
    s.title("最终判断标准", "项目长期可维护性的三个硬原则", dark=True)
    principles = [
        ("文件语义优先", "KiCad S-expression、AST 和 IR 先正确，画布只是结果。"),
        ("投影层可替换", "SVG、PixiJS、X6、DOM 都应能从同一 IR 派生。"),
        ("验证成为契约", "round-trip、编辑后重解析、SVG 快照、像素回归、KiCad CLI。"),
    ]
    for idx, (head, body) in enumerate(principles):
        y = 1.75 + idx * 1.25
        color = [PALETTE["gold"], PALETTE["pcb"], PALETTE["copper"]][idx]
        s.line(1.05, y + 0.08, 1.05, y + 0.78, color, 4)
        s.text(1.35, y, 3.1, 0.3, head, 23, PALETTE["white"], True)
        s.text(4.45, y + 0.08, 6.8, 0.35, body, 15, "C9D1CD")
    s.rect(7.25, 5.92, 4.35, 0.5, PALETTE["kicad"], rounded=True)
    s.text(7.5, 6.1, 3.85, 0.12, "一句话：IR 是真相，Renderer 是投影，验证是契约。", 9.5, PALETTE["white"], True, align="center")
    s.footer(18, dark=True)

    return prs, slides


def render_outputs(prs: Presentation, slides: list[SlideBuilder]) -> None:
    PREVIEW_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT_PPTX)

    outline_lines = ["# KiCad 符号 Web 渲染实现原理 PPT 大纲", ""]
    for idx, slide in enumerate(slides, 1):
        preview_path = PREVIEW_DIR / f"kicad-symbol-rendering-slide-{idx:02d}.png"
        slide.render_preview(preview_path)
        outline_lines.append(f"## {idx:02d}. {slide.meta.title}")
        text_items = [element.args[4] for element in slide.elements if element.kind == "text"]
        compact = []
        for item in text_items:
            item = re.sub(r"\s+", " ", item).strip()
            if item and item not in compact:
                compact.append(item)
        outline_lines.extend(f"- {item}" for item in compact[:8])
        outline_lines.append("")

    OUTLINE_FILE.write_text("\n".join(outline_lines), encoding="utf-8")
    create_montage(len(slides))
    write_qa(slides)


def create_montage(slide_count: int) -> None:
    thumbs: list[Image.Image] = []
    for idx in range(1, slide_count + 1):
        path = PREVIEW_DIR / f"kicad-symbol-rendering-slide-{idx:02d}.png"
        img = Image.open(path)
        img.thumbnail((360, 203))
        canvas = Image.new("RGB", (380, 232), pil_rgb(PALETTE["paper"]))
        canvas.paste(img, (10, 10))
        draw = ImageDraw.Draw(canvas)
        draw.text((12, 214), f"{idx:02d}", font=load_font(8), fill=pil_rgb(PALETTE["muted"]))
        thumbs.append(canvas)
    cols = 3
    rows = (slide_count + cols - 1) // cols
    montage = Image.new("RGB", (cols * 380, rows * 232), pil_rgb("E9E1D2"))
    for idx, thumb in enumerate(thumbs):
        montage.paste(thumb, ((idx % cols) * 380, (idx // cols) * 232))
    montage.save(MONTAGE_FILE)


def write_qa(slides: list[SlideBuilder]) -> None:
    required = [
        "[Content_Types].xml",
        "ppt/presentation.xml",
        "ppt/slides/slide1.xml",
        f"ppt/slides/slide{len(slides)}.xml",
    ]
    with zipfile.ZipFile(OUTPUT_PPTX) as archive:
        names = set(archive.namelist())
        missing = [name for name in required if name not in names]

    preview_files = sorted(PREVIEW_DIR.glob("kicad-symbol-rendering-slide-*.png"))
    quicklook_preview = PREVIEW_DIR / "quicklook" / f"{OUTPUT_PPTX.name}.png"
    if quicklook_preview.exists():
        preview_note = (
            "逐页 PNG 由同一几何描述的 PIL 预览通道生成，用于检查排版节奏和明显重叠；"
            f"另检测到 macOS Quick Look 对保存后 PPTX 生成的封面缩略图 `{quicklook_preview.relative_to(ROOT)}`。"
            "当前环境缺少 LibreOffice/pdftoppm，未做全量 PowerPoint 渲染器级逐页截图。"
        )
    else:
        preview_note = (
            "逐页 PNG 由同一几何描述的 PIL 预览通道生成，用于检查排版节奏和明显重叠；"
            "当前环境缺少 LibreOffice/pdftoppm，未做 PowerPoint 渲染器级逐页截图。"
        )
    qa = [
        "# PPT QA 记录",
        "",
        f"- 源文档：`{SOURCE_DOC.relative_to(ROOT)}`",
        f"- PPTX：`{OUTPUT_PPTX.relative_to(ROOT)}`",
        f"- 幻灯片数量：{len(slides)}",
        f"- PNG 预览数量：{len(preview_files)}",
        f"- PPTX 包结构：{'通过' if not missing else '缺失 ' + ', '.join(missing)}",
        f"- 文件大小：{OUTPUT_PPTX.stat().st_size} bytes",
        f"- 预览说明：{preview_note}",
        "",
        "## 自检结论",
        "",
        "- 页面均有标题、主视觉或结构化图形，不是纯文字页。",
        "- PPTX 使用原生文本框、形状、线条和表格式形状组合，保留可编辑性。",
        "- 内容覆盖目标、格式基础、分层架构、解析链路、IR、布局、坐标、PixiJS、SVG、导出、安全、验证、当前状态和演进路线。",
    ]
    QA_FILE.write_text("\n".join(qa), encoding="utf-8")


def main() -> None:
    prs, slides = create_deck()
    render_outputs(prs, slides)
    print(f"created {OUTPUT_PPTX}")
    print(f"slides {len(slides)}")
    print(f"previews {PREVIEW_DIR}")


if __name__ == "__main__":
    main()
